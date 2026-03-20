import os
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_report(output_path):
    document = Document()
    
    # Title
    title = document.add_heading('Comprehensive Project Report', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph('House Price Prediction: Aura Estates')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Date
    from datetime import datetime
    date_p = document.add_paragraph(datetime.now().strftime('%Y-%m-%d'))
    date_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Executive Summary
    document.add_heading('1. Executive Summary', level=1)
    document.add_paragraph(
        "This report outlines the end-to-end Machine Learning pipeline developed for predicting house prices using the Ames Housing dataset. "
        "The project encompasses data preprocessing, model selection and training, rigorous evaluation, and a FastAPI-based web application "
        "designed for real-time predictions. The best-performing model developed is an XGBoost Regressor with a Root Mean Squared Error (RMSE) of approximately 26,164."
    )

    # Project Architecture
    document.add_heading('2. Project Architecture', level=1)
    document.add_paragraph(
        "The system architecture is designed to support a robust flow from raw datasets to an interactive user interface, consisting of the following stages: "
    )
    stages = [
        "Raw Data Acquisition: Ingesting train and test sets from the Ames Housing context.",
        "Data Preprocessing: Handling missing values, scaling, and categorical encoding.",
        "Model Training & Tuning: Comparing base models (Linear Regression, Ridge, RandomForest) and optimizing the selected champion (XGBoost) using Optuna.",
        "Artifact Storage: Saving robust preprocessors and ML models.",
        "Web Deployment: Utilizing a FastAPI routing backend exposed to end-users for generating predictions on-the-fly."
    ]
    for stage in stages:
        document.add_paragraph(stage, style='List Bullet')

    # Data Description
    document.add_heading('3. Data Description & Features', level=1)
    document.add_paragraph(
        "The raw dataset includes 79 features detailing various aspects of residential homes in Ames, Iowa. The target variable to predict is `SalePrice`. "
        "For efficiency and relevance in the web interface, the deployment prioritizes several high-impact features, including:"
    )
    features = [
        "GrLivArea: Above grade (ground) living area square feet.",
        "OverallQual: Rates the overall material and finish of the house (Scale 1-10).",
        "TotalBsmtSF: Total square feet of basement area.",
        "FullBath: Full bathrooms above grade.",
        "GarageCars: Size of garage in car capacity.",
        "YearBuilt: Original construction date."
    ]
    for feature in features:
        document.add_paragraph(feature, style='List Bullet')

    # Data Preprocessing Strategy
    document.add_heading('4. Preprocessing Strategy', level=1)
    document.add_paragraph(
        "Extensive cleaning and transformation procedures were established in `preprocessing.py`. Highlights include:"
    )
    document.add_paragraph("- Missing Value Imputation: Numeric variables are filled with the mean, while categorical ones are filled with 'none' (assuming absence like no pool/fence).", style='List Bullet')
    document.add_paragraph("- Feature Scaling: Variables are bounded to the [0, 1] interval utilizing MinMaxScaler for model stability.", style='List Bullet')
    document.add_paragraph("- One-Hot Encoding: Used to transform categorical strings into binary inputs, expanding the dataset to 286 informative dimensions.", style='List Bullet')
    document.add_paragraph("- Dimensionality Restraint: The `Id` feature and columns with over 50% missing values (e.g., Alley, PoolQC) are completely dropped.", style='List Bullet')

    # Machine Learning and Modeling
    document.add_heading('5. Modeling and Evaluation', level=1)
    document.add_paragraph(
        "A rigorous approach was employed to select the most predictive model. Several architectures were trained and evaluated on their Root Mean Squared Error (RMSE):"
    )
    
    # Adding a table for results
    table = document.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Model'
    hdr_cells[1].text = 'Status'
    hdr_cells[2].text = 'Performance (RMSE)'

    records = (
        ('XGBoost', 'Winner', '~26,164'),
        ('RandomForest', 'Top Contender', '~27,032'),
        ('Linear Regression', 'Baseline', '~28,136'),
        ('Ridge Regression', 'Evaluated', '~28,500+')
    )
    for model, status, rmse in records:
        row_cells = table.add_row().cells
        row_cells[0].text = model
        row_cells[1].text = status
        row_cells[2].text = rmse

    document.add_paragraph(
        "\nThe final XGBoost model underwent hyperparameter tuning via Optuna. Parameters refined include `learning_rate`, `max_depth`, and `n_estimators`. "
        "Post-training, the model and preprocessing scaler/encoder were saved as Joblib artifacts (`xgboost_model.joblib` and `transformers.joblib`)."
    )

    # Web Deployment Using FastAPI
    document.add_heading('6. Deploying the FastAPI Web Application', level=1)
    document.add_paragraph(
        "The project transcends analysis by offering real-time endpoint capabilities constructed via FastAPI. Elements of the deployment include: "
    )
    document.add_paragraph("- Strict Validation: Pydantic schemas enforce bounds on user input (e.g., YearBuilt must be valid history up to 2026).", style='List Bullet')
    document.add_paragraph("- Automated Inference: Live inputs traverse the exact preprocessing pipeline utilized during training prior to evaluation.", style='List Bullet')
    document.add_paragraph("- REST Interface: A simple UI communicates POST requests to `/predict` and receives JSON-encapsulated housing price predictions.", style='List Bullet')

    # Conclusion
    document.add_heading('7. Conclusion', level=1)
    document.add_paragraph(
        "The House Price Prediction application successfully demonstrates a sophisticated pipeline moving from raw Iowa real estate characteristics to a robust, "
        "deployed API endpoint. The meticulous methodology regarding preprocessing, algorithm selection, and hyperparameter tuning has resulted in an inference model "
        "that provides highly dependable pricing estimations."
    )

    # Save Document
    document.save(output_path)
    print(f"Report comprehensively generated at: {output_path}")

if __name__ == "__main__":
    report_path = os.path.join(os.getcwd(), 'Aura_Estates_Project_Report.docx')
    create_report(report_path)
